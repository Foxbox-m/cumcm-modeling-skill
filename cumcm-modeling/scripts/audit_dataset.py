#!/usr/bin/env python3
"""Audit CUMCM CSV/XLSX data with explicit evidence boundaries.

The script exposes deterministic structure facts separately from heuristic
risks and questions that still need a problem-specific interpretation. It
never labels a field as leaked or cleans a value automatically.
"""

from __future__ import annotations

import argparse
import csv
import math
import re
import statistics
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any


MISSING_TEXT = {"", "na", "n/a", "null", "none", "nan", "-"}
DATE_PATTERNS = (
    "%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S",
    "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M",
)
ID_NAME_PATTERN = re.compile(r"(?:^|[_\-\s])(?:id|code|no|编号|序号|编码|代码)(?:$|[_\-\s])", re.I)
HEADER_HINT_PATTERN = re.compile(r"日期|时间|编号|序号|名称|类别|单位|总计|合计|数量|金额|地点|地区")


def is_missing(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    return isinstance(value, str) and value.strip().lower() in MISSING_TEXT


def _normalise_headers(raw: list[Any], width: int) -> list[str]:
    headers = [(str(cell).strip() if cell is not None else "") for cell in raw]
    headers = [cell or f"column_{index + 1}" for index, cell in enumerate(headers)]
    headers.extend(f"column_{index + 1}" for index in range(len(headers), width))
    return headers


def _validate_header_row(header_row: int) -> None:
    if header_row < 1:
        raise ValueError("--header-row 必须是从 1 开始的正整数")


def read_csv(path: Path, header_row: int = 1) -> tuple[list[str], list[list[Any]]]:
    _validate_header_row(header_row)
    last_error: UnicodeDecodeError | None = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                data = list(csv.reader(handle))
            break
        except UnicodeDecodeError as exc:
            last_error = exc
    else:
        raise ValueError(f"无法按 UTF-8 或 GB18030 解码：{last_error}")
    if len(data) < header_row:
        raise ValueError(f"数据不足以使用第 {header_row} 行作为表头")
    width = max(len(row) for row in data)
    headers = _normalise_headers(data[header_row - 1], width)
    rows = [row + [""] * (width - len(row)) for row in data[header_row:]]
    return headers, rows


def read_xlsx(path: Path, sheet: str | None, header_row: int = 1) -> tuple[list[str], list[list[Any]]]:
    _validate_header_row(header_row)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("读取 XLSX 需要 openpyxl；可先另存为 CSV") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet and sheet not in workbook.sheetnames:
            raise ValueError(f"工作表不存在：{sheet}；可选：{', '.join(workbook.sheetnames)}")
        worksheet = workbook[sheet] if sheet else workbook.active
        data = [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    if len(data) < header_row:
        raise ValueError(f"工作表不足以使用第 {header_row} 行作为表头")
    width = max(len(row) for row in data)
    headers = _normalise_headers(data[header_row - 1], width)
    rows = [row + [None] * (width - len(row)) for row in data[header_row:]]
    return headers, rows


def load_table(path: Path, sheet: str | None, header_row: int = 1) -> tuple[list[str], list[list[Any]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv(path, header_row)
    if suffix == ".xlsx":
        return read_xlsx(path, sheet, header_row)
    if suffix == ".xls":
        raise ValueError("旧版 XLS 暂不直接读取，请在表格软件中另存为 XLSX 或 CSV")
    raise ValueError("仅支持 CSV、XLSX；旧版 XLS 请先转换")


def as_number(value: Any) -> float | None:
    if is_missing(value) or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    scale = 0.01 if text.endswith("%") else 1.0
    if text.endswith("%"):
        text = text[:-1]
    try:
        return float(text) * scale
    except ValueError:
        return None


def parse_temporal(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if is_missing(value):
        return None
    text = str(value).strip()
    for pattern in DATE_PATTERNS:
        try:
            return datetime.strptime(text, pattern)
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        return None


def likely_date_sequence(values: list[Any]) -> tuple[bool, bool]:
    parsed = [parse_temporal(value) for value in values if not is_missing(value)]
    if len(parsed) < 2 or any(value is None for value in parsed):
        return False, False
    parsed_dates = [value for value in parsed if value is not None]
    return True, all(left <= right for left, right in zip(parsed_dates, parsed_dates[1:]))


def _column_index(headers: list[str], name: str | None, option: str) -> int | None:
    if not name:
        return None
    if name not in headers:
        raise ValueError(f"{option} 不存在：{name}；可选：{', '.join(headers)}")
    return headers.index(name)


def inspect_xlsx(path: Path, sheet: str | None) -> tuple[list[str], list[str], list[str]]:
    """Return deterministic workbook facts, heuristic header clues, and questions."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("--inspect 读取 XLSX 需要 openpyxl；可先另存为 CSV") from exc
    workbook = load_workbook(path, read_only=False, data_only=False)
    evidence: list[str] = []
    risks: list[str] = []
    questions: list[str] = []
    try:
        selected = [sheet] if sheet else list(workbook.sheetnames)
        for name in selected:
            if name not in workbook.sheetnames:
                raise ValueError(f"工作表不存在：{name}；可选：{', '.join(workbook.sheetnames)}")
            worksheet = workbook[name]
            hidden_rows = [str(index) for index, dimension in worksheet.row_dimensions.items() if dimension.hidden]
            hidden_cols = [str(index) for index, dimension in worksheet.column_dimensions.items() if dimension.hidden]
            merged = [str(region) for region in worksheet.merged_cells.ranges]
            formulas = 0
            comments = 0
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                        formulas += 1
                    if cell.comment is not None:
                        comments += 1
            evidence.append(
                f"工作表 {name!r}：{worksheet.max_row} 行 × {worksheet.max_column} 列；"
                f"合并区域 {len(merged)}；隐藏行 {len(hidden_rows)}；隐藏列 {len(hidden_cols)}；"
                f"公式单元格 {formulas}；批注单元格 {comments}"
            )
            if merged:
                evidence.append(f"  合并区域示例：{', '.join(merged[:8])}")
            if hidden_rows or hidden_cols:
                evidence.append(
                    f"  隐藏维度：行 {', '.join(hidden_rows[:12]) or '无'}；列 {', '.join(hidden_cols[:12]) or '无'}"
                )

            for row_index in range(1, min(worksheet.max_row, 12) + 1):
                values = [worksheet.cell(row_index, column).value for column in range(1, worksheet.max_column + 1)]
                present = [value for value in values if not is_missing(value)]
                if not present:
                    continue
                text_count = sum(isinstance(value, str) for value in present)
                hints = [str(value).strip() for value in present if HEADER_HINT_PATTERN.search(str(value))]
                score = len(hints) + (1 if len(set(map(str, present))) == len(present) else 0)
                if score or row_index == 1:
                    risks.append(
                        f"工作表 {name!r} 第 {row_index} 行可能是表头线索："
                        f"非空 {len(present)}，文本占比 {text_count / len(present):.0%}，"
                        f"关键词 {', '.join(hints[:6]) or '无'}"
                    )
            if merged or hidden_rows or hidden_cols or formulas or comments:
                questions.append(f"工作表 {name!r} 含特殊结构；请确认合并、隐藏、公式和批注是否属于建模数据口径。")
        if len(workbook.sheetnames) > 1 and sheet is None:
            questions.append("工作簿包含多个工作表；请明确指定 --sheet，避免误用活动工作表。")
    finally:
        workbook.close()
    return evidence, risks, questions


def audit(
    path: Path,
    sheet: str | None,
    inspect: bool = False,
    header_row: int = 1,
    time_column: str | None = None,
    target_column: str | None = None,
    prediction_cutoff: str | None = None,
) -> int:
    headers, rows = load_table(path, sheet, header_row)
    evidence: list[str] = [f"数据集：{path}", f"规模：{len(rows)} 行 × {len(headers)} 列", f"采用第 {header_row} 行作为表头"]
    risks: list[str] = []
    questions: list[str] = []
    if inspect:
        if path.suffix.lower() == ".xlsx":
            workbook_evidence, workbook_risks, workbook_questions = inspect_xlsx(path, sheet)
            evidence.extend(workbook_evidence)
            risks.extend(workbook_risks)
            questions.extend(workbook_questions)
        else:
            questions.append("--inspect 仅对 XLSX 提供工作簿侦察；当前 CSV 仅执行表格审计。")

    if len(set(headers)) != len(headers):
        duplicates = [name for name, count in Counter(headers).items() if count > 1]
        evidence.append(f"重复列名：{', '.join(duplicates)}")
        questions.append("请确认重复列名对应的实际含义，并在建模前显式重命名。")
    else:
        evidence.append("重复列名：未发现")

    normalized_rows = [tuple("" if is_missing(value) else str(value).strip() for value in row) for row in rows]
    evidence.append(f"重复数据行：{len(normalized_rows) - len(set(normalized_rows))}")
    missing_counts = [sum(is_missing(row[index]) for row in rows) for index in range(len(headers))]
    missing_items = [f"{name}: {count} ({count / max(len(rows), 1):.1%})" for name, count in zip(headers, missing_counts) if count]
    evidence.append("缺失值：" + ("；".join(missing_items) if missing_items else "未发现"))

    numeric_columns: dict[str, list[float]] = {}
    for index, name in enumerate(headers):
        present = [row[index] for row in rows if not is_missing(row[index])]
        numbers = [as_number(value) for value in present]
        if present and all(value is not None for value in numbers):
            values = [float(value) for value in numbers if value is not None]
            numeric_columns[name] = values
            std = statistics.stdev(values) if len(values) > 1 else 0.0
            evidence.append(
                f"数值列 {name}: n={len(values)}, min={min(values):.6g}, median={statistics.median(values):.6g}, "
                f"mean={statistics.fmean(values):.6g}, std={std:.6g}, max={max(values):.6g}"
            )
            unique_ratio = len(set(values)) / len(values)
            if len(values) >= 3 and (len(set(values)) == 1 or unique_ratio <= 0.05):
                risks.append(f"数值列 {name} 近常量候选：非缺失值唯一率 {unique_ratio:.1%}；请结合题意判断是否有信息量。")
    if not numeric_columns:
        evidence.append("数值列：未识别出纯数值列")

    for index, name in enumerate(headers):
        present = [row[index] for row in rows if not is_missing(row[index])]
        if len(present) < 3:
            continue
        unique_ratio = len({str(value).strip() for value in present}) / len(present)
        name_hint = bool(ID_NAME_PATTERN.search(name))
        parsed_numbers = [as_number(value) for value in present]
        integer_like = all(value is not None and float(value).is_integer() for value in parsed_numbers)
        if unique_ratio >= 0.8:
            risks.append(f"列 {name} 高基数候选：非缺失唯一率 {unique_ratio:.1%}；不等于应删除或属于泄漏。")
        if name_hint or (unique_ratio >= 0.95 and integer_like):
            risks.append(f"列 {name} 疑似标识列：名称线索={'是' if name_hint else '否'}，唯一率 {unique_ratio:.1%}；请确认是否仅作索引。")

    if len(numeric_columns) >= 3 and rows:
        indices = [headers.index(name) for name in numeric_columns]
        sums = []
        for row in rows:
            values = [as_number(row[index]) for index in indices]
            if all(value is not None for value in values):
                sums.append(sum(float(value) for value in values if value is not None))
        if sums:
            mean_sum = statistics.fmean(sums)
            spread = statistics.pstdev(sums) if len(sums) > 1 else 0.0
            if (abs(mean_sum - 1.0) <= 0.03 and spread <= 0.03) or (abs(mean_sum - 100.0) <= 3.0 and spread <= 3.0):
                risks.append(f"多列数值行和稳定接近 {mean_sum:.4g}，可能是成分数据；请检查闭合效应和零值处理。")

    selected_time_index = _column_index(headers, time_column, "--time-column")
    date_columns: list[str] = []
    for index, name in enumerate(headers):
        is_date, monotonic = likely_date_sequence([row[index] for row in rows])
        if is_date:
            date_columns.append(name)
            if monotonic:
                evidence.append(f"时间列 {name}：按当前行序单调")
            else:
                risks.append(f"时间列 {name}：按当前行序不单调；请检查排序、重复时间戳和跨表拼接。")
    if selected_time_index is not None and headers[selected_time_index] not in date_columns:
        questions.append(f"显式指定的时间列 {headers[selected_time_index]!r} 未能按常见日期格式解析；请确认格式或单位。")
    elif selected_time_index is None and len(date_columns) > 1:
        questions.append(f"自动识别出多个日期列（{', '.join(date_columns)}）；请用 --time-column 指定建模时间轴。")

    target_index = _column_index(headers, target_column, "--target-column")
    if target_index is not None and prediction_cutoff is None:
        questions.append(f"已指定目标列 {target_column!r}，但没有 --prediction-cutoff；无法判断预测边界，请补充时间语义。")
    if prediction_cutoff is not None:
        if selected_time_index is None:
            questions.append("提供了 --prediction-cutoff，但未提供可解析的 --time-column；无法进行时间边界核对。")
        else:
            time_values = [parse_temporal(row[selected_time_index]) for row in rows]
            cutoff_time = parse_temporal(prediction_cutoff)
            if cutoff_time is None:
                cutoff_number = as_number(prediction_cutoff)
                numeric_times = [as_number(row[selected_time_index]) for row in rows]
                nonmissing_times = [value for value in numeric_times if value is not None]
                if cutoff_number is not None and len(nonmissing_times) == sum(not is_missing(row[selected_time_index]) for row in rows):
                    after_cutoff = [index for index, value in enumerate(numeric_times) if value is not None and value > cutoff_number]
                else:
                    questions.append(f"预测截止点 {prediction_cutoff!r} 无法按时间或数值解析。")
                    after_cutoff = []
            else:
                after_cutoff = [index for index, value in enumerate(time_values) if value is not None and value > cutoff_time]
            evidence.append(f"按显式预测截止点 {prediction_cutoff!r}，截止点后的行数：{len(after_cutoff)}")
            if target_index is not None:
                populated_target_after = sum(not is_missing(rows[index][target_index]) for index in after_cutoff)
                if populated_target_after:
                    questions.append(
                        f"目标列 {target_column!r} 在预测截止点后仍有 {populated_target_after} 个非空值；"
                        "这是结构事实，不自动判定信息泄漏，请确认该列是否会作为预测输入。"
                    )

    print(f"数据集：{path}")
    print("\n确定性结构证据：")
    for item in evidence:
        print(f"  - {item}")
    print("\n启发式风险候选：")
    for item in risks or ["未发现；这不等于不存在语义风险。"]:
        print(f"  - {item}")
    print("\n待解释问题：")
    for item in questions or ["暂无；仍需结合题意、单位和建模目标复核。"]:
        print(f"  - {item}")
    print("\n审计边界：脚本只报告结构事实、启发式风险和待解释问题；不自动清洗数据，不断言目标泄漏，也不替代题意与量纲判断。")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="审计 CUMCM 赛题 CSV/XLSX 数据")
    parser.add_argument("dataset", type=Path)
    parser.add_argument("--sheet", help="XLSX 工作表名称；默认活动工作表")
    parser.add_argument("--header-row", type=int, default=1, help="表头所在行（从 1 开始，默认 1）")
    parser.add_argument("--time-column", help="显式指定时间列名，用于时间顺序与预测边界核对")
    parser.add_argument("--target-column", help="显式指定目标列名；不会据此自动判定泄漏")
    parser.add_argument("--prediction-cutoff", help="显式指定预测截止点（常用日期或数值格式）")
    parser.add_argument("--inspect", action="store_true", help="侦察 XLSX 工作表、合并/隐藏结构、公式/批注和表头线索")
    args = parser.parse_args()
    if not args.dataset.is_file():
        print(f"[错误] 文件不存在：{args.dataset}", file=sys.stderr)
        return 2
    try:
        return audit(args.dataset, args.sheet, inspect=args.inspect, header_row=args.header_row,
                     time_column=args.time_column, target_column=args.target_column,
                     prediction_cutoff=args.prediction_cutoff)
    except (OSError, ValueError, RuntimeError) as exc:
        print(f"[错误] {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
