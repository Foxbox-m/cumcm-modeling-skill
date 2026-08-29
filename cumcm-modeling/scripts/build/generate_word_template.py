#!/usr/bin/env python3
"""Generate the editable CUMCM Word template from explicit OOXML tokens."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Iterable

try:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor
except ImportError as exc:  # pragma: no cover - depends on the caller's environment
    if __name__ == "__main__":
        print("[错误] 生成 Word 模板需要 python-docx；请安装后重试。", file=sys.stderr)
        raise SystemExit(2) from exc
    raise


GRAY = RGBColor(128, 128, 128)
BLACK = RGBColor(0, 0, 0)
CHINESE_FONT = "SimSun"
HEADING_FONT = "SimHei"
LATIN_FONT = "Times New Roman"


def set_run_font(run, name: str = CHINESE_FONT, size: float = 12, *, bold: bool = False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), LATIN_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def style_font(style, name: str, size: float, *, bold: bool = False):
    style.font.name = LATIN_FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), LATIN_FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = BLACK
    # Keep the hierarchy neutral even when the document inherits a themed
    # accent from Word's default template.
    color = style._element.get_or_add_rPr().find(qn("w:color"))
    if color is not None:
        color.attrib.pop(qn("w:themeColor"), None)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            tag = "w:" + edge
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            for key, value in kwargs[edge].items():
                element.set(qn("w:" + key), str(value))


def set_table_width(table, widths: Iterable[int]):
    widths = list(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_field(paragraph, instruction: str, placeholder: str = "1"):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = placeholder
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def add_omml_formula(paragraph):
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    for token in ("x", "i", "=", "y", "i", "+", "b"):
        run = OxmlElement("m:r")
        text = OxmlElement("m:t")
        text.text = token
        run.append(text)
        omath.append(run)
    omath_para.append(omath)
    paragraph._p.append(omath_para)


def add_placeholder(doc, text: str, *, indent: bool = True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.first_line_indent = Cm(0.74) if indent else Cm(0)
    run = p.add_run(text)
    set_run_font(run, size=12, color=GRAY)
    return p


def add_hidden_placeholder(doc, text: str, *, indent: bool = False):
    """Add an instructional prompt hidden in the default Word view."""
    paragraph = add_placeholder(doc, text, indent=indent)
    for run in paragraph.runs:
        r_pr = run._element.get_or_add_rPr()
        if r_pr.find(qn("w:vanish")) is None:
            r_pr.append(OxmlElement("w:vanish"))
    return paragraph


def add_bookmark_start(paragraph, name: str, bookmark_id: int):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    paragraph._p.append(start)


def add_bookmark_end(paragraph, bookmark_id: int):
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.append(end)


def add_caption(doc, label: str, text: str, *, bookmark_name: str | None = None, bookmark_id: int = 1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(label + " ")
    set_run_font(run, size=10.5, bold=True, color=BLACK)
    if bookmark_name:
        add_bookmark_start(p, bookmark_name, bookmark_id)
    add_field(p, f"SEQ {label}", "1")
    field = p._p.find(qn("w:fldSimple"))
    if field is not None:
        field_run = field.find(qn("w:r"))
        if field_run is not None:
            field_r_pr = field_run.find(qn("w:rPr"))
            if field_r_pr is None:
                field_r_pr = OxmlElement("w:rPr")
                field_run.insert(0, field_r_pr)
            if field_r_pr.find(qn("w:b")) is None:
                field_r_pr.append(OxmlElement("w:b"))
    if bookmark_name:
        add_bookmark_end(p, bookmark_id)
    run = p.add_run("：" + text)
    set_run_font(run, size=10.5, bold=False, color=BLACK)
    return p


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    # CUMCM profile is single-column as a Skill internal gate.  Keep one
    # section and make the OOXML setting explicit for downstream audits.
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), "1")
    normal = doc.styles["Normal"]
    style_font(normal, CHINESE_FONT, 12)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.space_after = Pt(0)
    for name, size, before, after in (("Heading 1", 14, 10, 6), ("Heading 2", 12, 8, 4), ("Heading 3", 12, 4, 2)):
        style = doc.styles[name]
        style_font(style, HEADING_FONT, size, bold=True)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        if name == "Heading 1":
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title = doc.styles["Title"]
    style_font(title, HEADING_FONT, 16, bold=True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("第 ")
    set_run_font(run, size=10.5)
    add_field(footer, "PAGE", "1")
    run = footer.add_run(" 页")
    set_run_font(run, size=10.5)


def chinese_question(index: int) -> str:
    digits = "一二三四五六七八九十"
    return digits[index - 1] if 1 <= index <= len(digits) else str(index)


def appendix_label(index: int) -> str:
    value = index
    label = ""
    while value:
        value, remainder = divmod(value - 1, 26)
        label = chr(ord("A") + remainder) + label
    return label


def add_semantic_task(doc, index: int):
    """Add one question-aware task chapter with removable functional prompts."""
    question = chinese_question(index)
    doc.add_heading(f"5.{index} 问题{question}的模型建立与求解（按题面替换）", level=2)
    add_placeholder(doc, "填写：最短充分直觉引线、目标/输入、模型、求解、正式结果、结果现象/机制、必要局部验证与直接作答。")
    add_placeholder(doc, "复核：真实 Cascade/Parallel/Hybrid（级联/并列/混合）关系、基线缺口、活跃约束、局部简化依据与实现一致性。")
    for suffix, prompt in (
        (1, "模型建立：请替换为当前题的数学对象、关系或约束标题"),
        (2, "求解与计算结果：请替换为当前题的求解动作或结果标题"),
        (3, "结果分析与局部检验：请替换为当前题的证据或边界标题"),
    ):
        doc.add_heading(f"5.{index}.{suffix} {prompt}", level=3)
        add_placeholder(doc, "从现象或短直觉立即进入数学对象、关系或约束；纯解析、极短或真实级联任务可合并功能层，但模型/推导、结果、局部证据和直接答案仍须可定位。")
    add_hidden_placeholder(
        doc,
        "详细写作提示（默认隐藏，不构成固定小节）：Fast Narrative Lead 先给最短充分现象/直觉；Phase A 对每个 2.N/5.N 一次完成“现象→直觉→数学→求解→正式结果→现象/机制→必要局部验证→直接作答”，仅在风险尚未关闭时保留足以关闭主要风险的必要局部验证，若不同主要风险需要不同证据，分别保留；局部验证/不确定性按上述风险口径处理；随后 Phase B 只做一次全文读取；完整论文或最终交付时，Competition Paper Voice 与 Claim Calibration 不以用户额外点名为前提，但仍仅处理实际观察到的问题；Caveat/限定语也仅处理本次全文扫描发现的问题，所有维度均不机械执行无问题项；局部任务只归类当前任务范围内实际观察到的问题。合并清单后一次局部修复并只复验受影响维度；Semantic Emphasis 仅在发现具体语义导航缺口，或最终正式交付确需扫读导航时启用，局部语态/限定语清理不得自动增加或调整粗体。自然完成“目标与输出→输入与来源→数学关系/数学结构/模型→推导与求解→量化结果→机制解释与局部验证→直接作答”，并说明选择理由、变量/关系/目标/约束、参数/假设/单位/可辨识性和适用条件。标题方面，题面有显式小问时保留并替换问题编号与语义标题；合并小问只为避免重复，仍须保留对应标题号；题面没有显式小问时，使用对象/机制/决策/验证形成的语义标题，不虚构问题编号。多问先判断 Cascade/Parallel/Hybrid：只有真实级联才写传递关系与增量，并列任务写独立对象/场景/困难/表示/验证出口，混合任务只写真实共享核心与接口。主模型要展开结构困难与关系来源、相对最小基线所需的机制及参数/约束来源与主导因素；求解/离散设置交代初值/随机性、停止准则、容差、成本与失败/降级路径；结果部分先提出图表要回答的问题，再按数据结构选择图型并读取具体证据（有序轨迹才用折线；变量关系可用散点/拟合与残差；分布可用直方图/KDE/ECDF/箱线；矩阵可用带标签热图；物理机理图需标注当前题坐标/边界/变量；标明接口、分支、终止的流程图）；结果按“结果→题目现象/状态→主导变量/约束/几何/机理→题意”解释，局部简化依据留在首次依赖处；最后用 1—2 句直接回答题面并写清下一任务接口。Phase B 只改可见语言并保护公式、label、cite key、图表路径、正式数值、单位、冻结参数、硬约束和 LaTeX 结构；若数字或公式需要改变，回到事实源。复核基线归因、活跃约束解释、声明模型与实际实现一致、离散/步长/收敛依据。",
    )
    add_hidden_placeholder(doc, "内部工程追踪字段只留在项目记录和支撑材料，不进入摘要与正文。", indent=False)


def build_document(output: Path, *, questions: int = 1):
    if questions <= 0:
        raise ValueError("questions 必须是正整数")
    doc = Document()
    configure_document(doc)
    title = doc.add_paragraph(style="Title")
    title.add_run("论文题目：方法/问题对象")
    for run in title.runs:
        set_run_font(run, HEADING_FONT, 16, bold=True, color=BLACK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("摘    要")
    set_run_font(run, HEADING_FONT, 14, bold=True, color=BLACK)
    add_placeholder(doc, "总体段：填写研究对象、数据/机理基础、总体目标和模型体系。")
    for index in range(1, questions + 1):
        question = chinese_question(index)
        add_placeholder(doc, f"针对问题{question}：按困难→模型/关键求解→正式结果→一句现象/机制组织，只保留一项关键验证，并给出可核对结果。")
    add_hidden_placeholder(
        doc,
        "摘要写作提示（默认隐藏）：摘要第一页且原则上一页，采用总体段加逐问独立段；每问按困难→模型/关键求解→正式结果→一句现象/机制组织，只保留一项关键验证；四至五问摘要明显过短只提示复核证据，不按字数自动扩写。",
    )
    p = doc.add_paragraph()
    run = p.add_run("关键词：请替换为对象词；核心模型词；求解/验证词")
    set_run_font(run, size=12, color=GRAY)
    doc.add_page_break()

    doc.add_heading("一、问题重述", level=1)
    doc.add_heading("1.1 问题背景", level=2)
    add_placeholder(doc, "用自己的话概括现实对象、决策目标或观测任务，不机械复制题面。")
    doc.add_heading("1.2 问题要求", level=2)
    for index in range(1, questions + 1):
        add_placeholder(doc, f"问题{chinese_question(index)}：只重述本问输入、输出、硬约束和交付量，不提前给出模型或答案。", indent=False)
    doc.add_heading("二、问题分析", level=1)
    for index in range(1, questions + 1):
        question = chinese_question(index)
        doc.add_heading(f"2.{index} 问题{question}的分析（按题面替换）", level=2)
        add_placeholder(doc, f"第一自然句直接说明问题{question}的对象、动作或输出；先判断 Cascade/Parallel/Hybrid（级联/并列/混合），再按真实关系说明困难、数学表示、求解路线以及结果与验证出口，不能只列算法名。")
    doc.add_heading("三、模型假设", level=1)
    add_placeholder(doc, "集中列出全局共享且影响模型成立性的假设，并说明简化影响；局部假设在首次使用处解释。")
    doc.add_heading("四、符号说明", level=1)
    table = doc.add_table(rows=2, cols=3)
    set_table_width(table, [1800, 5160, 2200])
    for cell, text in zip(table.rows[0].cells, ("符号", "含义", "单位")):
        cell.text = text
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=10.5, bold=True, color=BLACK)
    for cell, text in zip(table.rows[1].cells, ("请替换", "请替换为物理含义", "请替换")):
        cell.text = text
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=10.5, color=GRAY)
    for cell in table.rows[0].cells:
        set_cell_border(cell, top={"val": "single", "sz": "8", "color": "000000"}, bottom={"val": "single", "sz": "4", "color": "000000"})
    for cell in table.rows[-1].cells:
        set_cell_border(cell, bottom={"val": "single", "sz": "8", "color": "000000"})

    doc.add_heading("五、模型的建立与求解", level=1)
    add_hidden_placeholder(
        doc,
        "共享核心准入（默认隐藏）：只有同时服务至少两问、能被后续直接调用、且前置定义可减少重复并不引入单问细节时，才在第五章开头定义；否则随对应 5.N 局部建立。",
    )
    for index in range(1, questions + 1):
        add_semantic_task(doc, index)

    doc.add_heading("六、模型的分析与检验", level=1)
    add_placeholder(doc, "只汇总跨问题、系统级或共同风险的检验；逐问局部验证仍紧邻 5.N 的相应结果。")
    doc.add_heading("七、模型的评价", level=1)
    doc.add_heading("7.1 模型的优点", level=2)
    add_placeholder(doc, "优点必须回指已展示的模型结构、结果或验证证据。")
    doc.add_heading("7.2 模型的缺点与改进", level=2)
    add_placeholder(doc, "先集中列出由结果暴露的整体局限、触发条件与影响方向，再统一提出可执行改进；公式成立所需的局部简化依据留在首次依赖处，避免重复免责。")
    doc.add_heading("AI 工具使用声明", level=1)
    add_placeholder(doc, "当前 2026 规则要求所有参赛队在参考文献之前如实声明；未来届次先复核规则。该声明不作为编号章节。")
    doc.add_heading("参考文献", level=1)
    add_placeholder(doc, "按实际核验过的来源填写作者、题名、来源、年份和 DOI/URL。")

    doc.add_page_break()
    doc.add_heading("附录", level=1)
    for index in range(1, questions + 1):
        question = chinese_question(index)
        doc.add_heading(f"附录 {appendix_label(index)}：问题{question}求解程序及说明", level=2)
        add_placeholder(doc, f"先说明本问特异核心函数/函数签名和主求解逻辑，并交代对应正文 5.{index} 节的输入、输出和功能；再纳入完整执行源文件；文件索引不能替代代码。")
    doc.add_heading(f"附录 {appendix_label(questions + 1)}：公共函数与共享模块", level=2)
    add_placeholder(doc, "在各问特异核心逻辑之后只后置一次共享完整代码；列各问题实际调用的公共函数和共享模块，并说明题号到入口脚本的对应关系。")
    doc.add_heading(f"附录 {appendix_label(questions + 2)}：提交支撑材料文件清单", level=2)
    add_placeholder(doc, "只列真正提交的题目代码、共享函数、结果文件和必要数据；不列项目简报、内部账本、检查日志、QA 预览或内部清单。")

    add_hidden_placeholder(
        doc,
        "图表与公式指导（默认隐藏）：插入图或表后添加中文题注并使用 Word 交叉引用更新域。三线表应设置明确列宽、表头和单位；图件正文说明应检查文字、单位、图例和灰度可辨性。一个独立公式/等式链承担一个论证动作并单独成行；公式前写依据或目的，公式后紧跟符号、单位、参数来源和后续作用。多个约束、分段状态或联立方程用左大括号分行，关键公式编号并交叉引用。",
        indent=False,
    )
    core = doc.core_properties
    core.author = ""
    core.last_modified_by = ""
    core.comments = ""
    core.subject = ""
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description="生成可编辑 CUMCM Word 论文模板")
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parents[2] / "assets/paper-template/cumcm_template.docx")
    parser.add_argument("--questions", type=int, default=1, help="题面显式问题数，用于展开摘要、1.2、2.N、5.N 和附录")
    args = parser.parse_args()
    if args.questions <= 0:
        parser.error("--questions 必须是正整数")
    build_document(args.output, questions=args.questions)
    print(f"已生成 Word 模板：{args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
