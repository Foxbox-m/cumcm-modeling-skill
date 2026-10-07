from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.comments import Comment


SCRIPT = Path(__file__).parents[2] / "cumcm-modeling" / "scripts" / "audit_dataset.py"
SPEC = importlib.util.spec_from_file_location("audit_dataset", SCRIPT)
assert SPEC and SPEC.loader
audit_dataset = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(audit_dataset)


def test_inspect_irregular_xlsx_reports_structure_and_semantic_boundary(tmp_path, capsys):
    path = tmp_path / "irregular.xlsx"
    workbook = Workbook()
    note = workbook.active
    note.title = "说明"
    note["A1"] = "数据说明"
    sheet = workbook.create_sheet("数据")
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "原始观测"
    sheet["A2"] = "日期"
    sheet["B2"] = "编号"
    sheet["C2"] = "值"
    sheet["D2"] = "目标"
    rows = [
        ("2024-01-02", 1, "=1+2", 10),
        ("2024-01-01", 2, 3, 11),
        ("2024-01-03", 3, 3, 12),
    ]
    for row in rows:
        sheet.append(row)
    sheet["C3"].comment = Comment("computed", "fixture")
    sheet.row_dimensions[4].hidden = True
    sheet.column_dimensions["D"].hidden = True
    workbook.save(path)

    result = audit_dataset.audit(
        path,
        "数据",
        inspect=True,
        header_row=2,
        time_column="日期",
        target_column="目标",
        prediction_cutoff="2024-01-01",
    )
    output = capsys.readouterr().out
    assert result == 0
    assert "确定性结构证据" in output
    assert "启发式风险候选" in output
    assert "待解释问题" in output
    assert "合并区域 1" in output
    assert "隐藏行 1" in output
    assert "隐藏列 1" in output
    assert "公式单元格 1" in output
    assert "批注单元格 1" in output
    assert "时间列 日期：按当前行序不单调" in output
    assert "目标列 '目标' 在预测截止点后仍有" in output
    assert "不自动判定信息泄漏" in output
    assert "不自动清洗数据" in output


def test_csv_risks_are_separated_from_questions(tmp_path, capsys):
    path = tmp_path / "risk.csv"
    path.write_text(
        "日期,编号,常量,目标\n"
        "2024-01-02,1,7,3\n"
        "2024-01-01,2,7,4\n"
        "2024-01-03,3,7,5\n",
        encoding="utf-8",
    )
    assert audit_dataset.audit(path, None) == 0
    output = capsys.readouterr().out
    assert "启发式风险候选" in output
    assert "列 编号 疑似标识列" in output
    assert "数值列 常量 近常量候选" in output
    assert "时间列 日期：按当前行序不单调" in output
    assert "判定为泄漏" not in output


def test_invalid_semantic_column_is_a_clear_error(tmp_path, capsys):
    path = tmp_path / "data.csv"
    path.write_text("a,b\n1,2\n", encoding="utf-8")
    with pytest.raises(ValueError, match="--time-column 不存在"):
        audit_dataset.audit(path, None, time_column="missing")
